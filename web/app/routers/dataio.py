"""
데이터 입출력 라우터 (dataio.py)
================================
- CSV 주소록 가져오기: 연락처 대량 등록
- 엑셀 내보내기: 발송 이력·통계를 .xlsx로

엑셀은 openpyxl 없이도 동작하도록 설계:
- openpyxl 있으면 .xlsx 생성
- 없으면 CSV(UTF-8 BOM)로 대체 (엑셀에서 바로 열림)
표준 라이브러리(csv, io) 우선.
"""
import csv
import io
from datetime import datetime

from fastapi import APIRouter, Depends, UploadFile, File, Query
from fastapi.responses import StreamingResponse

from ..core import db, errors
from ..deps import get_current_user, require_manager

router = APIRouter(prefix='/api/data', tags=['data-io'])


# ==================== CSV 주소록 가져오기 ====================
# 허용 헤더 별칭 (유연한 매핑)
_NAME_KEYS = {'name', '이름', '성명', '담당자', 'contact', 'contact_name'}
_FAX_KEYS = {'fax', 'fax_number', '팩스', '팩스번호', 'number', '번호', 'faxno'}
_COMPANY_KEYS = {'company', '회사', '업체', '거래처', 'org', 'organization'}
_MEMO_KEYS = {'memo', '메모', 'note', '비고'}


def _detect_columns(headers):
    """헤더에서 이름·번호·회사·메모 컬럼 인덱스 탐지."""
    idx = {'name': None, 'fax': None, 'company': None, 'memo': None}
    for i, h in enumerate(headers):
        key = (h or '').strip().lstrip('\ufeff').lower()
        if idx['name'] is None and key in _NAME_KEYS:
            idx['name'] = i
        elif idx['fax'] is None and key in _FAX_KEYS:
            idx['fax'] = i
        elif idx['company'] is None and key in _COMPANY_KEYS:
            idx['company'] = i
        elif idx['memo'] is None and key in _MEMO_KEYS:
            idx['memo'] = i
    return idx


def _read_csv_bytes(raw):
    """바이트를 CSV 행 리스트로. 인코딩 자동 감지(utf-8/cp949)."""
    for enc in ('utf-8-sig', 'utf-8', 'cp949', 'euc-kr'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise errors.ValidationError(
            detail='CSV 인코딩을 인식할 수 없습니다',
            user_message='CSV 파일 인코딩을 확인해주세요 (UTF-8 또는 CP949).')
    # 구분자 자동 감지 (콤마/탭/세미콜론)
    sample = text[:2000]
    delim = ','
    for d in (',', '\t', ';'):
        if sample.count(d) >= sample.count(delim):
            delim = d
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    return [row for row in reader if any(c.strip() for c in row)]


@router.post('/contacts/import/preview')
async def import_preview(file: UploadFile = File(...), user=Depends(get_current_user)):
    """
    CSV 미리보기 - 실제 등록 전에 컬럼 매핑·행 수 확인.
    """
    raw = await file.read()
    if not raw:
        raise errors.EmptyFile()
    rows = _read_csv_bytes(raw)
    if len(rows) < 1:
        raise errors.ValidationError(user_message='CSV에 데이터가 없습니다.')
    headers = rows[0]
    idx = _detect_columns(headers)
    has_header = idx['name'] is not None or idx['fax'] is not None
    # 헤더 없으면 첫 컬럼=이름, 둘째=번호로 가정
    if not has_header:
        idx = {'name': 0, 'fax': 1 if len(headers) > 1 else None,
               'company': 2 if len(headers) > 2 else None, 'memo': None}
    data_rows = rows[1:] if has_header else rows
    # 미리보기 (최대 5행)
    preview = []
    for r in data_rows[:5]:
        preview.append({
            'name': r[idx['name']].strip() if idx['name'] is not None and idx['name'] < len(r) else '',
            'fax_number': r[idx['fax']].strip() if idx['fax'] is not None and idx['fax'] < len(r) else '',
            'company': r[idx['company']].strip() if idx['company'] is not None and idx['company'] < len(r) else '',
        })
    return {
        'headers': headers, 'detected': idx, 'has_header': has_header,
        'total_rows': len(data_rows), 'preview': preview,
    }


@router.post('/contacts/import')
async def import_contacts(
    file: UploadFile = File(...),
    is_shared: bool = Query(default=False),
    skip_duplicates: bool = Query(default=True),
    user=Depends(get_current_user)):
    """
    CSV에서 연락처 대량 등록.
    - 이름·번호 없는 행은 건너뜀
    - skip_duplicates: 같은 번호가 이미 있으면 건너뜀
    """
    raw = await file.read()
    if not raw:
        raise errors.EmptyFile()
    rows = _read_csv_bytes(raw)
    if not rows:
        raise errors.ValidationError(user_message='CSV에 데이터가 없습니다.')
    headers = rows[0]
    idx = _detect_columns(headers)
    has_header = idx['name'] is not None or idx['fax'] is not None
    if not has_header:
        idx = {'name': 0, 'fax': 1 if len(headers) > 1 else None,
               'company': 2 if len(headers) > 2 else None, 'memo': None}
    data_rows = rows[1:] if has_header else rows

    # 기존 번호 (중복 체크용)
    existing = set()
    if skip_duplicates:
        rows_db = db.query(
            'SELECT fax_number FROM contacts WHERE owner_id=%s OR is_shared=true',
            (user['user_id'],))
        existing = {r['fax_number'].strip() for r in rows_db if r.get('fax_number')}

    added, skipped, errored = 0, 0, 0
    error_samples = []
    for r in data_rows:
        try:
            name = r[idx['name']].strip() if idx['name'] is not None and idx['name'] < len(r) else ''
            fax = r[idx['fax']].strip() if idx['fax'] is not None and idx['fax'] < len(r) else ''
            company = r[idx['company']].strip() if idx['company'] is not None and idx['company'] < len(r) else None
            memo = r[idx['memo']].strip() if idx['memo'] is not None and idx['memo'] < len(r) else None
            if not name or not fax:
                skipped += 1
                continue
            if skip_duplicates and fax in existing:
                skipped += 1
                continue
            db.execute(
                'INSERT INTO contacts(owner_id, name, fax_number, company, memo, is_shared) '
                'VALUES(%s,%s,%s,%s,%s,%s)',
                (user['user_id'], name, fax, company, memo, is_shared))
            existing.add(fax)
            added += 1
        except Exception as e:
            errored += 1
            if len(error_samples) < 3:
                error_samples.append(str(e)[:80])
    return {
        'ok': True, 'added': added, 'skipped': skipped, 'errored': errored,
        'message': f'{added}건 추가, {skipped}건 건너뜀' + (f', {errored}건 오류' if errored else ''),
        'error_samples': error_samples,
    }


@router.get('/contacts/template')
def contacts_template(user=Depends(get_current_user)):
    """가져오기용 CSV 템플릿 다운로드."""
    buf = io.StringIO()
    buf.write('\ufeff')  # UTF-8 BOM (엑셀 한글)
    w = csv.writer(buf)
    w.writerow(['이름', '팩스번호', '회사', '메모'])
    w.writerow(['홍길동', '02-1234-5678', '○○상사', ''])
    w.writerow(['김철수', '031-987-6543', '△△기업', ''])
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]), media_type='text/csv',
        headers={'Content-Disposition': 'attachment; filename="contacts_template.csv"'})


# ==================== 엑셀/CSV 내보내기 ====================
def _try_xlsx(sheet_name, headers, rows):
    """openpyxl로 xlsx 생성 시도. 없으면 None 반환."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    # 헤더 스타일
    header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val)
    # 열 너비 자동
    for c, h in enumerate(headers, 1):
        maxlen = max([len(str(h))] + [len(str(row[c-1])) for row in rows if c-1 < len(row)] + [8])
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = min(maxlen + 4, 40)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _csv_fallback(headers, rows):
    """CSV(UTF-8 BOM) 생성 - 엑셀에서 바로 열림."""
    buf = io.StringIO()
    buf.write('\ufeff')
    w = csv.writer(buf)
    w.writerow(headers)
    for row in rows:
        w.writerow(row)
    return io.BytesIO(buf.getvalue().encode('utf-8'))


def _export_response(basename, headers, rows, ascii_name='export'):
    """xlsx 우선, 없으면 csv로 다운로드 응답.
    파일명 한글은 HTTP 헤더(latin-1)에 직접 못 넣으므로,
    ASCII 파일명 + RFC5987(filename*) 방식으로 한글명 병행 제공."""
    from urllib.parse import quote
    xlsx = _try_xlsx(basename[:30], headers, rows)
    ts = datetime.now().strftime('%Y%m%d')
    ext = 'xlsx' if xlsx else 'csv'
    # 한글 파일명은 UTF-8 퍼센트 인코딩 (RFC5987)
    kr_name = quote(f'{basename}_{ts}.{ext}')
    ascii_fallback = f'{ascii_name}_{ts}.{ext}'
    disposition = f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{kr_name}"
    if xlsx:
        return StreamingResponse(
            xlsx, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': disposition})
    csv_buf = _csv_fallback(headers, rows)
    return StreamingResponse(
        csv_buf, media_type='text/csv',
        headers={'Content-Disposition': disposition})


_STATUS_KR = {'sent': '완료', 'failed': '실패', 'queued': '대기', 'pending': '준비',
              'sending': '전송중', 'retry_wait': '재시도대기', 'cancelled': '취소'}


@router.get('/export/send-history')
def export_send_history(
    status: str = None, days: int = Query(default=90, le=365),
    user=Depends(get_current_user)):
    """발송 이력 엑셀 내보내기."""
    cond = ['created_at >= CURRENT_DATE - %s::int']
    params = [days]
    if user['role'] not in ('admin', 'manager'):
        cond.append('user_id=%s'); params.append(user['user_id'])
    if status:
        cond.append('status=%s'); params.append(status)
    sql = ('SELECT created_at, to_number, to_name, file_name, pages, status, '
           'cover_used, retry_count FROM fax_jobs WHERE ' + ' AND '.join(cond) +
           ' ORDER BY id DESC')
    data = db.query(sql, tuple(params))
    headers = ['발송시각', '받는번호', '받는사람', '파일명', '페이지', '상태', '표지', '재시도']
    rows = [[
        str(d['created_at'])[:19], d['to_number'], d.get('to_name') or '',
        d.get('file_name') or '', d.get('pages') or 0,
        _STATUS_KR.get(d['status'], d['status']),
        '예' if d.get('cover_used') else '', d.get('retry_count') or 0,
    ] for d in data]
    return _export_response('발송이력', headers, rows, ascii_name='send_history')


@router.get('/export/received')
def export_received(days: int = Query(default=90, le=365), user=Depends(get_current_user)):
    """수신 이력 엑셀 내보내기."""
    cond = ['received_at >= CURRENT_DATE - %s::int']
    params = [days]
    if user['role'] not in ('admin', 'manager'):
        cond.append('(user_id=%s OR user_id IS NULL)'); params.append(user['user_id'])
    sql = ('SELECT received_at, from_number, to_number, pages, is_read, memo '
           'FROM fax_received WHERE ' + ' AND '.join(cond) + ' ORDER BY id DESC')
    data = db.query(sql, tuple(params))
    headers = ['수신시각', '발신번호', '수신번호', '페이지', '읽음', '메모']
    rows = [[
        str(d['received_at'])[:19], d.get('from_number') or '', d.get('to_number') or '',
        d.get('pages') or 0, '읽음' if d.get('is_read') else '안읽음', d.get('memo') or '',
    ] for d in data]
    return _export_response('수신이력', headers, rows, ascii_name='received')


@router.get('/export/stats')
def export_stats(user=Depends(require_manager)):
    """부서별 통계 엑셀 내보내기 (매니저+)."""
    data = db.query(
        "SELECT d.name, COUNT(j.id) AS sent_count, "
        "COUNT(j.id) FILTER (WHERE j.status='sent') AS sent_ok "
        "FROM departments d LEFT JOIN fax_jobs j ON j.dept_id=d.id "
        "GROUP BY d.id, d.name ORDER BY d.name")
    headers = ['부서', '발송건수', '성공건수', '성공률']
    rows = [[
        d['name'], d['sent_count'], d['sent_ok'],
        f"{round(100*d['sent_ok']/d['sent_count'],1)}%" if d['sent_count'] else '-',
    ] for d in data]
    return _export_response('부서별통계', headers, rows, ascii_name='dept_stats')


@router.get('/export/client-report')
def export_client_report(
    days: int = Query(default=30, le=365),
    to_number: str = None,
    user=Depends(require_manager)):
    """고객사별 발송 실적 리포트 (매니저+).
    번호별로 발송/성공/실패/페이지를 집계 - 고객 증빙용 명세.
    to_number 지정 시 해당 고객만."""
    cond = ['created_at >= CURRENT_DATE - %s::int']
    params = [days]
    if to_number:
        cond.append('to_number=%s'); params.append(to_number)
    sql = (
        "SELECT to_number, MAX(to_name) AS to_name, "
        "COUNT(*) AS total, "
        "COUNT(*) FILTER (WHERE status='sent') AS ok, "
        "COUNT(*) FILTER (WHERE status='failed') AS failed, "
        "COUNT(*) FILTER (WHERE status IN ('queued','pending','sending','retry_wait','scheduled')) AS pending, "
        "COALESCE(SUM(pages) FILTER (WHERE status='sent'),0) AS pages_sent, "
        "MIN(created_at) AS first_at, MAX(created_at) AS last_at "
        "FROM fax_jobs WHERE " + ' AND '.join(cond) +
        " GROUP BY to_number ORDER BY total DESC")
    data = db.query(sql, tuple(params))
    headers = ['받는번호', '고객명', '총건수', '성공', '실패', '진행중',
               '성공률', '전송페이지', '첫발송', '마지막발송']
    rows = [[
        d['to_number'], d.get('to_name') or '',
        d['total'], d['ok'], d['failed'], d['pending'],
        f"{round(100*d['ok']/d['total'],1)}%" if d['total'] else '-',
        d['pages_sent'],
        str(d['first_at'])[:10], str(d['last_at'])[:10],
    ] for d in data]
    return _export_response('고객사별_발송실적', headers, rows, ascii_name='client_report')
